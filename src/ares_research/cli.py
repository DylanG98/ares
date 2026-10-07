"""Operator and in-agent commands. No credentials are printed."""

from __future__ import annotations

import argparse
import json
import os
import sys
from dataclasses import asdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .auth import hermes_auth_status
from .domain import CompanyIdentity, ValuationInputs
from .extraction import extract_company_fact, extract_pdf
from .organization import OrganizationService
from .paperclip import PaperclipClient, PaperclipError
from .review import IndependentReviewer
from .smoke import run_smoke
from .spreadsheet import SpreadsheetRecalculator
from .storage import atomic_json
from .workbook import FinancialWorkbook
from .workflow import AnalysisWorkflow


def main():
    parser = argparse.ArgumentParser(
        prog="ares", description="Research y Valuación sobre Paperclip/Hermes"
    )
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("bootstrap")
    commands.add_parser("doctor")
    activate = commands.add_parser("activate")
    activate.add_argument("--model", required=True)
    hire = commands.add_parser("hire")
    hire.add_argument(
        "--role",
        required=True,
        choices=["research", "accounting", "business", "valuation", "review"],
    )
    hire.add_argument("--decision", type=Path, required=True)
    request = commands.add_parser("request")
    request.add_argument("file", type=Path)
    plan = commands.add_parser("plan")
    plan.add_argument("--case", required=True)
    plan.add_argument("--parent", required=True)
    smoke = commands.add_parser("smoke")
    smoke.add_argument("--output", type=Path, default=Path("artifacts/synthetic"))
    extract = commands.add_parser("extract-pdf")
    extract.add_argument("file", type=Path)
    xbrl = commands.add_parser("extract-xbrl")
    xbrl.add_argument("file", type=Path)
    for argument in ["taxonomy", "concept", "unit", "end", "cutoff"]:
        xbrl.add_argument(f"--{argument}", required=True)
    xbrl.add_argument("--start")
    model = commands.add_parser("model")
    model.add_argument("file", type=Path)
    model.add_argument("--history", type=Path, required=True)
    model.add_argument("--sources", type=Path, required=True)
    model.add_argument("--output", type=Path, required=True)
    model.add_argument("--synthetic", action="store_true")
    review = commands.add_parser("review")
    review.add_argument("--inputs", type=Path, required=True)
    review.add_argument("--workbook", type=Path, required=True)
    recalc = commands.add_parser("recalculate")
    recalc.add_argument("file", type=Path)
    recalc.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        if args.command == "model":
            inputs = ValuationInputs.model_validate_json(args.file.read_text())
            FinancialWorkbook().write(
                args.output,
                inputs,
                json.loads(args.history.read_text()),
                json.loads(args.sources.read_text()),
                synthetic=args.synthetic,
            )
            result = {"workbook": str(args.output), "status": "requires_recalculation_and_review"}
        elif args.command == "review":
            inputs = ValuationInputs.model_validate_json(args.inputs.read_text())
            reviewer = IndependentReviewer()
            findings = reviewer.inspect_workbook(args.workbook)
            book = load_workbook(args.workbook, data_only=True)
            observed = {
                book["Valuacion"].cell(row, 1).value: book["Valuacion"].cell(row, 8).value
                for row in range(3, 6)
            }
            book.close()
            findings.extend(reviewer.check_values(inputs, observed))
            result = {
                "findings": [asdict(f) for f in findings],
                "status": "blocked" if findings else "numeric_checks_passed_not_agent_signoff",
            }
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 1 if findings else 0
        elif args.command == "recalculate":
            result = {
                "workbook": str(SpreadsheetRecalculator().recalculate(args.file, args.output))
            }
        elif args.command == "smoke":
            result = run_smoke(args.output)
        elif args.command == "extract-pdf":
            result = extract_pdf(args.file)
        elif args.command == "extract-xbrl":
            result = extract_company_fact(
                json.loads(args.file.read_text()),
                args.taxonomy,
                args.concept,
                args.unit,
                date.fromisoformat(args.end),
                date.fromisoformat(args.cutoff),
                date.fromisoformat(args.start) if args.start else None,
            )
        else:
            root = Path(os.environ.get("ARES_ROOT", Path(__file__).resolve().parents[2]))
            data = Path(os.environ.get("ARES_DATA", root / ".runtime"))
            client = PaperclipClient.from_env()
            try:
                if not client.agent_mode:
                    client.board_login()
                org = OrganizationService(client, root, data)
                if args.command == "bootstrap":
                    result = org.bootstrap()
                elif args.command == "activate":
                    result = org.activate(args.model)
                elif args.command == "hire":
                    result = org.hire(args.role, args.decision)
                elif args.command == "request":
                    result = AnalysisWorkflow(org).request(
                        CompanyIdentity.model_validate_json(args.file.read_text())
                    )
                elif args.command == "plan":
                    result = AnalysisWorkflow(org).plan(args.case, args.parent)
                else:
                    state = org.state()
                    agents = client.request("GET", f"/api/companies/{state['company_id']}/agents")
                    runs = client.request(
                        "GET",
                        f"/api/companies/{state['company_id']}/heartbeat-runs?limit=1000&summary=true",
                    )
                    commissioning = client.request(
                        "GET", f"/api/issues/{state['commissioning_issue_id']}"
                    )
                    auth = hermes_auth_status(data)
                    successful = {r["agentId"] for r in runs if r["status"] == "succeeded"}
                    expected = {r["key"] for r in org.manifest["roles"]}
                    actual = {(a.get("metadata") or {}).get("ares_role") for a in agents}
                    readiness = {
                        "chatgpt_authenticated": bool(auth.get("logged_in")),
                        "six_distinct_roles": actual == expected and len(agents) == 6,
                        "all_agents_have_successful_run": bool(agents)
                        and all(a["id"] in successful for a in agents),
                        "commissioning_accepted_by_ceo": commissioning["status"] == "done",
                    }
                    result = {
                        "paperclip": client.request("GET", "/api/health"),
                        "organization": state,
                        "agents": [
                            {k: a.get(k) for k in ["id", "name", "status", "lastHeartbeatAt"]}
                            for a in agents
                        ],
                        "chatgpt": auth,
                        "expected_roles": 6,
                        "actual_roles": len(agents),
                        "readiness": readiness,
                        "operational": all(readiness.values()),
                        "qualification": "Successful runs do not replace review of commissioning evidence.",
                    }
                    atomic_json(data / "state/doctor.json", result)
            finally:
                client.close()
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        return 0
    except (ValueError, PermissionError, FileNotFoundError, PaperclipError, RuntimeError) as exc:
        print(f"ARES: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
