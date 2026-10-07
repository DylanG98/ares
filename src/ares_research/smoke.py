"""Technical acceptance using synthetic evidence. Never impersonates agent execution."""

from __future__ import annotations

import hashlib
from dataclasses import asdict
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

from .domain import CompanyIdentity, FinancialFact, Source, ValuationInputs
from .extraction import extract_company_fact, extract_pdf
from .finance import DcfModel
from .review import IndependentReviewer
from .storage import Dossier, atomic_json
from .workbook import FinancialWorkbook


def synthetic_inputs() -> ValuationInputs:
    scenarios = []
    for name, growth, margin, wacc, terminal, roic, narrative in [
        (
            "adverso",
            ["-.05", "0", ".02", ".02", ".02"],
            ".12",
            ".12",
            ".015",
            ".10",
            "Pérdida de volumen y presión competitiva reducen márgenes; recuperación lenta.",
        ),
        (
            "base",
            [".06", ".05", ".04", ".03", ".03"],
            ".18",
            ".10",
            ".02",
            ".15",
            "Crecimiento de volumen moderado y precios estables; reinversión disciplinada.",
        ),
        (
            "favorable",
            [".10", ".08", ".06", ".05", ".04"],
            ".22",
            ".09",
            ".025",
            ".18",
            "Mayor utilización de capacidad y expansión de segmento elevan rentabilidad.",
        ),
    ]:
        scenarios.append(
            dict(
                name=name,
                narrative=narrative,
                revenue_growth=growth,
                ebit_margin=margin,
                tax_rate=".25",
                da_to_sales=".03",
                capex_to_sales=".05",
                nwc_to_sales=".12",
                wacc=wacc,
                terminal_growth=terminal,
                terminal_roic=roic,
            )
        )
    return ValuationInputs(
        currency="USD",
        basis="nominal",
        rate_basis="nominal",
        sector="industrial",
        revenue=1000,
        debt=200,
        cash=50,
        non_operating_assets=10,
        minority_interest=5,
        diluted_shares=100,
        shares_per_instrument=1,
        lease_treatment="operating_in_cashflows",
        lease_liabilities=0,
        source_ids=["synthetic-annual"],
        scenarios=scenarios,
    )


def run_smoke(output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    identity = CompanyIdentity(
        legal_name="Empresa Sintética Ares SA",
        ticker="SYNTH",
        exchange="TEST",
        instrument="ordinary_share",
        currency="USD",
        cutoff_date=date(2025, 3, 1),
        accounting_standard="Synthetic US GAAP-like fixture",
        sector="industrial",
        synthetic=True,
    )
    dossier = Dossier(output / "dossiers", identity)
    pdf = output / "synthetic-original.pdf"
    if not pdf.exists():
        canvas = Canvas(str(pdf), invariant=1)
        canvas.drawString(50, 790, "SYNTHETIC TEST DATA - NOT A REAL COMPANY")
        canvas.drawString(50, 760, "Ares Synthetic Company: fiscal year 2024, currency USD")
        canvas.drawString(50, 730, "Revenue: 1000; Assets: 1500; Liabilities: 600; Equity: 900")
        canvas.save()
    content = pdf.read_bytes()
    source = Source(
        source_id="synthetic-annual",
        url="synthetic://ares/annual/2024",
        title="Synthetic annual fixture",
        publisher="Ares technical test",
        published_on=date(2025, 2, 1),
        retrieved_on=date(2025, 3, 1),
        sha256=hashlib.sha256(content).hexdigest(),
        kind="synthetic",
    )
    dossier.add_source(source, content)
    pages = extract_pdf(pdf)
    assert "Revenue: 1000" in pages[0]["text"]
    fact = FinancialFact(
        concept="Revenue",
        value=1000,
        unit="USD",
        currency="USD",
        entity=identity.legal_name,
        period_start=date(2024, 1, 1),
        period_end=date(2024, 12, 31),
        period_type="annual",
        source_id=source.source_id,
        locator="page 1: Revenue: 1000",
        original_verified=True,
    )
    dossier.validate_fact(fact)
    history = [
        {
            "year": year,
            "revenue": 800 + 50 * i,
            "ebit": (800 + 50 * i) * 0.18,
            "net_income": 80 + 5 * i,
            "cfo": 100 + 5 * i,
            "capex": 40 + 2.5 * i,
            "assets": 1300 + 50 * i,
            "liabilities": 520 + 20 * i,
            "equity": 780 + 30 * i,
            "source_id": "synthetic-annual",
        }
        for i, year in enumerate(range(2020, 2025))
    ]
    # Only 2024 revenue/balance are in the PDF; other historical cells are explicit fixture assumptions.
    historical_path = output / "synthetic-historicals.json"
    atomic_json(historical_path, {"synthetic": True, "rows": history})
    hist_source = Source(
        source_id="synthetic-history",
        url="synthetic://ares/history",
        title="Synthetic historical fixture",
        publisher="Ares technical test",
        published_on=date(2025, 2, 1),
        retrieved_on=date(2025, 3, 1),
        sha256=hashlib.sha256(historical_path.read_bytes()).hexdigest(),
        kind="synthetic",
    )
    dossier.add_source(hist_source, historical_path.read_bytes())
    for row in history:
        row["source_id"] = "synthetic-history"
    inputs = synthetic_inputs()
    results = {s.name: DcfModel(inputs).calculate(s) for s in inputs.scenarios}
    observed = {key: float(result.per_instrument) for key, result in results.items()}
    reviewer = IndependentReviewer()
    assert not reviewer.check_values(inputs, observed)
    corrupted = {**observed, "base": observed["base"] + 2}
    findings = reviewer.check_values(inputs, corrupted)
    assert any(f.code == "DCF_MISMATCH" and f.severity == "material" for f in findings)
    xlsx = output / "synthetic-model.xlsx"
    FinancialWorkbook().write(
        xlsx,
        inputs,
        history,
        [
            {**source.model_dump(mode="json"), "locator": "page 1"},
            {**hist_source.model_dump(mode="json"), "locator": "rows"},
        ],
        synthetic=True,
    )
    assert not reviewer.inspect_workbook(xlsx)
    wb = load_workbook(xlsx, data_only=True)
    assert wb.sheetnames == list(FinancialWorkbook.SHEETS)
    wb.close()
    report = Document()
    report.add_heading("ARES — PRUEBA SINTÉTICA", 0)
    report.add_paragraph(
        "No es un análisis de una empresa real. Prueba determinista de software; no acredita ejecuciones multiagente ni autenticación."
    )
    for name, value in observed.items():
        report.add_paragraph(f"{name}: {value:.6f} USD por instrumento (supuestos sintéticos).")
    report.add_paragraph(
        "Fuente: synthetic-original.pdf, página 1; históricos adicionales: synthetic-historicals.json."
    )
    report.save(output / "synthetic-report.docx")
    assert Document(output / "synthetic-report.docx").paragraphs
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "ARES — PRUEBA SINTÉTICA"
    slide.placeholders[
        1
    ].text = (
        "Artefacto técnico, no una conclusión de inversión ni una ejecución real de especialistas."
    )
    deck.save(output / "synthetic-summary.pptx")
    assert len(Presentation(output / "synthetic-summary.pptx").slides) == 1
    xbrl = {
        "entityName": identity.legal_name,
        "cik": "SYNTH",
        "facts": {
            "us-gaap": {
                "Revenues": {
                    "units": {
                        "USD": [
                            {
                                "start": "2024-01-01",
                                "end": "2024-12-31",
                                "filed": "2025-02-01",
                                "val": 1000,
                                "accn": "synthetic",
                            }
                        ]
                    }
                }
            }
        },
    }
    extracted = extract_company_fact(
        xbrl,
        "us-gaap",
        "Revenues",
        "USD",
        date(2024, 12, 31),
        identity.cutoff_date,
        date(2024, 1, 1),
    )
    assert extracted["value"] == 1000 and extracted["original_verified"] is False
    dossier.revise(
        "extracted", "research", [fact.model_dump(mode="json")], "Synthetic software extraction"
    )
    dossier.revise(
        "model", "valuation", inputs.model_dump(mode="json"), "Synthetic technical model"
    )
    dossier.revise(
        "reviews",
        "review",
        [asdict(f) for f in findings],
        "Deterministic injected-error check, not an agent review",
    )
    evidence = {
        "kind": "deterministic_technical_test",
        "synthetic": True,
        "case_id": identity.case_id,
        "pdf_extraction": "pass",
        "xbrl_extraction": "pass_requires_original_validation",
        "independent_numeric_check": "pass",
        "injected_error_detected": [asdict(f) for f in findings],
        "xlsx_formula_and_cache_inspection": "pass",
        "spreadsheet_engine_recalculation": "not_run",
        "docx_and_pptx_open": "pass",
        "valuations": observed,
        "real_agent_runs": "not_run_requires_chatgpt_login",
        "tokens": None,
        "cost": None,
    }
    atomic_json(output / "inputs.json", inputs.model_dump(mode="json"))
    atomic_json(output / "acceptance.json", evidence)
    return evidence
