"""Independent recomputation, deliberately not importing the valuation engine."""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .domain import Scenario, ValuationInputs


@dataclass(frozen=True)
class Finding:
    code: str
    severity: str
    owner: str
    message: str


class IndependentReviewer:
    def recalculate(self, inputs: ValuationInputs, scenario: Scenario) -> float:
        sales = float(inputs.revenue)
        flows = []
        margin, tax = float(scenario.ebit_margin), float(scenario.tax_rate)
        for growth in scenario.revenue_growth:
            prior = sales
            sales *= 1 + float(growth)
            flows.append(
                sales
                * (
                    margin * (1 - tax)
                    + float(scenario.da_to_sales)
                    - float(scenario.capex_to_sales)
                )
                - (sales - prior) * float(scenario.nwc_to_sales)
            )
        discount = float(scenario.wacc)
        growth = float(scenario.terminal_growth)
        steady_flow = (
            sales * margin * (1 - tax) * (1 + growth) * (1 - growth / float(scenario.terminal_roic))
        )
        terminal = steady_flow / (discount - growth)
        enterprise = math.fsum(
            flow * (1 + discount) ** (-year) for year, flow in enumerate(flows, 1)
        )
        enterprise += terminal * (1 + discount) ** (-len(flows))
        equity = (
            enterprise
            + float(inputs.cash)
            + float(inputs.non_operating_assets)
            - float(inputs.debt)
            - float(inputs.lease_liabilities)
            - float(inputs.minority_interest)
        )
        return equity / float(inputs.diluted_shares) * float(inputs.shares_per_instrument)

    def check_values(self, inputs: ValuationInputs, observed: dict[str, float]) -> list[Finding]:
        findings = []
        for scenario in inputs.scenarios:
            expected = self.recalculate(inputs, scenario)
            actual = observed.get(scenario.name)
            if (
                actual is None
                or not isinstance(actual, (int, float))
                or not math.isfinite(actual)
                or not math.isclose(actual, expected, rel_tol=1e-8, abs_tol=1e-6)
            ):
                findings.append(
                    Finding(
                        "DCF_MISMATCH",
                        "material",
                        "valuation",
                        f"{scenario.name}: observed {actual}; independent result {expected}",
                    )
                )
        return findings

    def inspect_workbook(self, path: Path) -> list[Finding]:
        formulas = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        findings = []
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        if "#REF!" in cell.value or "[" in cell.value:
                            findings.append(
                                Finding(
                                    "BROKEN_REFERENCE",
                                    "material",
                                    "valuation",
                                    f"{sheet.title}!{cell.coordinate}",
                                )
                            )
                        if cached[sheet.title][cell.coordinate].value is None:
                            findings.append(
                                Finding(
                                    "UNCALCULATED",
                                    "material",
                                    "valuation",
                                    f"{sheet.title}!{cell.coordinate}: recalculate in Excel/LibreOffice",
                                )
                            )
                    if cached[sheet.title][cell.coordinate].data_type == "e":
                        findings.append(
                            Finding(
                                "EXCEL_ERROR",
                                "material",
                                "valuation",
                                f"{sheet.title}!{cell.coordinate}",
                            )
                        )
        if "Controles" in cached:
            for label, value, *_ in cached["Controles"].iter_rows(values_only=True):
                if isinstance(label, str) and label.startswith("Balance "):
                    if (
                        not isinstance(value, (int, float))
                        or not math.isfinite(value)
                        or abs(value) > 1e-6
                    ):
                        findings.append(
                            Finding(
                                "BALANCE_NOT_RECONCILED",
                                "material",
                                "accounting",
                                f"{label}: {value}",
                            )
                        )
        formulas.close()
        cached.close()
        return findings

    @staticmethod
    def assert_publishable(
        author_id: str, reviewer_id: str, findings: list[Finding], evidence: list[str]
    ):
        if author_id == reviewer_id:
            raise PermissionError("Model author cannot approve their own work")
        if not evidence:
            raise ValueError("Review must identify persistent evidence")
        if any(f.severity in {"material", "critical"} for f in findings):
            raise ValueError("Unresolved material findings block publication")
