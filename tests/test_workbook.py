import pytest
from openpyxl import load_workbook

from ares_research.review import IndependentReviewer
from ares_research.smoke import synthetic_inputs
from ares_research.workbook import FinancialWorkbook


@pytest.mark.parametrize("equity", [None, 79])
def test_missing_or_unbalanced_history_blocks_review(tmp_path, equity):
    path = tmp_path / "model.xlsx"
    FinancialWorkbook().write(
        path,
        synthetic_inputs(),
        [
            {
                "year": 2024,
                "assets": 100,
                "liabilities": 20,
                "equity": equity,
            }
        ],
        [],
        synthetic=True,
    )
    reviewer = IndependentReviewer()
    assert any(f.code == "BALANCE_NOT_RECONCILED" for f in reviewer.inspect_workbook(path))
    if equity is None:
        book = load_workbook(path, data_only=True)
        assert book["Controles"]["B3"].value == "N/D"
        book.close()
